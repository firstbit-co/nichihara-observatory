"""
担当者が使うスプレッドシートのひな形(docs/sheet/nichihara-update.xlsx)を作る。
Google ドライブにアップロードすると Google スプレッドシートになる(使い方は docs/09-sheet.md)。

  python3 docs/sheet/make-template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / 'nichihara-update.xlsx'
HEAD = PatternFill('solid', fgColor='1F2B57')
NOTE = PatternFill('solid', fgColor='FFF4D6')
ROWS = 100


def sheet(wb, title, columns, sample, lists):
    ws = wb.create_sheet(title)
    for i, (name, width) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = width
    for i, v in enumerate(sample, start=1):
        c = ws.cell(row=2, column=i, value=v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    # 日付・時刻の列は「書式なしテキスト」にする。Google の CSV 書き出しは、列の型に合わないセルを空にすることがあるため
    for i, (name, _) in enumerate(columns, start=1):
        if name in ('日付', '開催日', '開始', '終了'):
            for r in range(2, ROWS + 1):
                ws.cell(row=r, column=i).number_format = '@'
    ws.freeze_panes = 'A2'
    names = [n for n, _ in columns]
    for col_name, options in lists.items():
        letter = ws.cell(row=1, column=names.index(col_name) + 1).column_letter
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
        dv.add(f'{letter}2:{letter}{ROWS}')
        ws.add_data_validation(dv)
    return ws


wb = Workbook()
wb.remove(wb.active)

sheet(
    wb,
    'お知らせ',
    [('公開', 12), ('日付', 13), ('種類', 14), ('タイトル', 36), ('本文', 70), ('重要', 8)],
    ['下書き', '2026/11/01', 'お知らせ', '(例)11月の観望会のご案内', '(例)本文をここに書きます。\n\n空行を入れると段落が分かれます。', ''],
    {'公開': ['公開する', '下書き'], '種類': ['お知らせ', 'イベント', '天文情報', '活動報告', '休館・営業', '道路・アクセス', 'メディア'], '重要': ['○']},
)
sheet(
    wb,
    'イベント',
    [('公開', 12), ('開催日', 13), ('開始', 9), ('終了', 9), ('種類', 14), ('イベント名', 30), ('ひとことで', 40), ('料金', 18), ('本文', 60)],
    ['下書き', '2026/12/14', '19:00', '22:00', '天文現象', '(例)ふたご座流星群の観望会', '(例)流れ星をさがします。', '通常の観望料金', '(例)暖かい服装でお越しください。'],
    {'公開': ['公開する', '下書き'], '種類': ['観望会', '天文現象', 'ワークショップ', '季節イベント', 'その他']},
)

how = wb.create_sheet('使い方', 0)
how.column_dimensions['A'].width = 100
lines = [
    ('日原天文台サイト お知らせ・イベントの更新', True),
    ('', False),
    ('1. 「お知らせ」または「イベント」のシートに、1行に1件ずつ書きます。', False),
    ('2. 「公開」を「公開する」にすると、30分ほどでサイトに出ます。「下書き」の行は出ません。', False),
    ('3. 直したいときは、その行を書きかえます。消したいときは「公開」を「下書き」にするか、行を削除します。', False),
    ('4. 日付は 2026/11/01 のように書きます。時刻は 19:00 のように書きます。', False),
    ('5. 「重要」に ○ を付けたお知らせは、サイトの全ページの上に帯で出ます(臨時休館など急ぎのときだけ)。', False),
    ('', False),
    ('書き方のきまり', True),
    ('・休館は毎週火曜・水曜と元日。臨時休館(金曜休館など)は、種類を「休館・営業」にして書きます。', False),
    ('・予約は受け付けていません。予約・定員・先着順は書きません。', False),
    ('・電話での確認はすすめません。料金が変わるとは書きません。', False),
    ('・「隣の」「日本初」は使いません。', False),
    ('きまりに合わない行はサイトに出ません。理由は確認ページで見られます:', False),
    ('https://firstbit-co.github.io/nichihara-observatory/admin/status.html', False),
    ('', False),
    ('1行目(見出し)とシートの名前は変えないでください。サイトが読めなくなります。', True),
]
for i, (text, bold) in enumerate(lines, start=1):
    c = how.cell(row=i, column=1, value=text)
    c.font = Font(bold=bold, size=13 if i == 1 else 11)
    if bold and i > 1:
        c.fill = NOTE

wb.save(OUT)
print('wrote', OUT)
